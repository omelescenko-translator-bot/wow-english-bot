import os
import io
import time
import openpyxl
try:
    from gtts import gTTS
except ImportError:
    gTTS = None
from aiogram import Router, types, F, Bot
from aiogram.types import (
    InlineKeyboardMarkup, InlineKeyboardButton, BufferedInputFile, WebAppInfo
)
from database.db import (
    add_card, get_all_cards, delete_card, get_stats, get_or_create_user,
    set_card_status, SUPPORTED_LANGS
)
from seed_initial_data import categorize
from translator import smart_translate
import config

router = Router()

USER_TRAIN_STATE = {} # user_id -> {'queue': [...], 'index': 0, 'session_count': 0}

def get_update_keyboard(user_id: int, card_id: int = None, is_single: bool = True, added_count: int = 1):
    webapp_url = config.get_webapp_url()
    v_timestamp = int(time.time())
    target_url = f"{webapp_url}?v={v_timestamp}&uid={user_id}" if (webapp_url and webapp_url.startswith("https://")) else webapp_url
    
    rows = []
    if card_id and is_single:
        rows.append([
            InlineKeyboardButton(text="🎧 Послушать", callback_data=f"tts_{card_id}"),
            InlineKeyboardButton(text="🗑️ Удалить", callback_data=f"del_{card_id}")
        ])
        
    if target_url and target_url.startswith("https://"):
        btn_label = f"🔄 📱 Обновить Mini App (+{added_count} фраз)" if added_count > 1 else "🔄 📱 Обновить Mini App"
        rows.append([
            InlineKeyboardButton(
                text=btn_label,
                web_app=WebAppInfo(url=target_url)
            )
        ])
    else:
        rows.append([
            InlineKeyboardButton(
                text="🔄 📱 Обновить бот и словарь",
                callback_data="reload_bot"
            )
        ])
        
    rows.append([
        InlineKeyboardButton(
            text="🃏 Начать тренировку карточек",
            callback_data="start_train_cb"
        )
    ])
    
    return InlineKeyboardMarkup(inline_keyboard=rows)

def parse_line_to_pair(line: str):
    line = line.strip()
    if not line or line.startswith('#'):
        return None, None
        
    for sep in ['\t', ' — ', ' - ', ':']:
        if sep in line:
            parts = line.split(sep, 1)
            en = parts[0].strip()
            ru = parts[1].strip()
            if any(c in 'абвгдеёжзийклмнопрстуфхцчшщъыьэюя' for c in en.lower()) and not any(c in 'абвгдеёжзийклмнопрстуфхцчшщъыьэюя' for c in ru.lower()):
                en, ru = ru, en
            return en, ru
            
    return line, ""

# --- FLASHCARD IN-CHAT TRAINING ---

@router.message(F.text == "🃏 Тренировка карточек")
async def start_flashcard_session(message: types.Message):
    user_id = message.from_user.id
    # Always random shuffle
    cards = get_all_cards(user_id, shuffle=True)
    if not cards:
        await message.answer(
            "📭 В твоём словаре пока нет карточек.\n"
            "Пришли мне фразу в формате:\n`phrase - перевод`\nчтобы добавить первую карточку!",
            parse_mode="Markdown"
        )
        return
        
    USER_TRAIN_STATE[user_id] = {
        'queue': cards,
        'index': 0,
        'session_count': 0
    }
    await send_train_card(message, user_id, edit=False)

async def send_train_card(event_target, user_id: int, edit: bool = False):
    state = USER_TRAIN_STATE.get(user_id)
    if not state:
        return
        
    queue = state['queue']
    index = state['index']
    session_count = state['session_count']
    
    if index >= len(queue):
        text = (
            "🎉 **Отличная работа! Тренировка завершена!**\n\n"
            f"Вы повторили все `{len(queue)}` карточек в этой сессии.\n"
            "Карточки распределены по статусам «Учу» и «Знаю»!"
        )
        if edit:
            await event_target.message.edit_text(text, parse_mode="Markdown")
        else:
            await event_target.answer(text, parse_mode="Markdown")
        return
        
    card = queue[index]
    total = len(queue)
    
    # 4 out of 5 cards RU -> EN, 1 out of 5 EN -> RU
    is_en_to_ru = (session_count % 5 == 4)
    
    if is_en_to_ru:
        mode_str = "🇬🇧 ➔ 🇷🇺 (1 из 5)"
        front_text = card['phrase_en']
        front_lbl = "ENGLISH 🇬🇧"
        tts_card_id = card['id']
    else:
        mode_str = "🇷🇺 ➔ 🇬🇧"
        front_text = card['phrase_ru']
        front_lbl = "РУССКИЙ 🇷🇺"
        tts_card_id = card['id']
        
    status_icon = "✅ Знаю" if card.get('status') == 'known' else "⏳ Учу"
    
    text = (
        f"🃏 **Карточка {index + 1} из {total}** | {mode_str}\n"
        f"📌 *Текущий статус: {status_icon}*\n\n"
        f"🏷️ **{front_lbl}:**\n"
        f"**{front_text}**\n\n"
        f"*(Нажмите кнопку ниже, чтобы увидеть перевод)*"
    )
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="🔄 Показать перевод", callback_data=f"flip_{card['id']}_{index}"),
            InlineKeyboardButton(text="🎧 Озвучить", callback_data=f"tts_{tts_card_id}")
        ]
    ])
    
    if edit:
        await event_target.message.edit_text(text, reply_markup=kb, parse_mode="Markdown")
    else:
        await event_target.answer(text, reply_markup=kb, parse_mode="Markdown")

@router.callback_query(F.data.startswith("flip_"))
async def flip_card_callback(callback: types.CallbackQuery):
    parts = callback.data.split("_")
    card_id = int(parts[1])
    index = int(parts[2])
    
    user_id = callback.from_user.id
    state = USER_TRAIN_STATE.get(user_id)
    if not state or index >= len(state['queue']):
        await callback.answer()
        return
        
    queue = state['queue']
    card = queue[index]
    total = len(queue)
    session_count = state['session_count']
    
    is_en_to_ru = (session_count % 5 == 4)
    mode_str = "🇬🇧 ➔ 🇷🇺 (1 из 5)" if is_en_to_ru else "🇷🇺 ➔ 🇬🇧"
    
    text = (
        f"🃏 **Карточка {index + 1} из {total}** | {mode_str}\n\n"
        f"🇷🇺 **Русский:** {card['phrase_ru']}\n"
        f"🇬🇧 **English:** **{card['phrase_en']}**\n\n"
        f"Оцените ваш ответ:"
    )
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="⏳ Учу", callback_data=f"status_learning_{index}"),
            InlineKeyboardButton(text="🎧 Озвучить", callback_data=f"tts_{card['id']}"),
            InlineKeyboardButton(text="✅ Знаю", callback_data=f"status_known_{index}")
        ]
    ])
    
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="Markdown")
    await callback.answer()

@router.callback_query(F.data.startswith("status_"))
async def status_card_callback(callback: types.CallbackQuery):
    parts = callback.data.split("_")
    new_status = parts[1] # 'learning' or 'known'
    index = int(parts[2])
    
    user_id = callback.from_user.id
    state = USER_TRAIN_STATE.get(user_id)
    
    if state and index < len(state['queue']):
        card = state['queue'][index]
        set_card_status(card['id'], new_status)
        card['status'] = new_status
        
        # If learning, push to back
        if new_status == 'learning':
            state['queue'].append(card)
            
        state['index'] += 1
        state['session_count'] += 1
        
    await send_train_card(callback, user_id, edit=True)
    await callback.answer(f"Статус: {'Знаю ✅' if new_status == 'known' else 'Учу ⏳'}")

# --- VOCABULARY & EXPORT ---

@router.message(F.text == "📚 Мой словарь")
async def show_vocabulary(message: types.Message):
    user_id = message.from_user.id
    cards = get_all_cards(user_id)
    stats = get_stats(user_id)
    
    if not cards:
        await message.answer(
            "📭 Твой словарь пока пуст.\n\n"
            "Пришли мне фразу в формате:\n`phrase - перевод`\n"
            "или перешли сообщение из канала, или отправь Excel-файл!",
            parse_mode="Markdown"
        )
        return
        
    preview = "\n".join([f"• **{c['phrase_en']}** — _{c['phrase_ru']}_ [{'✅' if c.get('status') == 'known' else '⏳'}]" for c in cards[:8]])
    
    text = (
        f"📚 **Твой личный словарь** (всего `{stats['total']}` фраз):\n"
        f"⏳ В процессе (Учу): `{stats['learning']}` | ✅ Знаю: `{stats['known']}`\n\n"
        f"{preview}\n\n"
        f"💡 *Показаны последние добавленные фразы.*"
    )
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📥 Скачать словарь в Excel", callback_data="export_excel")],
        [InlineKeyboardButton(text="🃏 Начать тренировку", callback_data="start_train_cb")]
    ])
    
    await message.answer(text, reply_markup=kb, parse_mode="Markdown")

@router.callback_query(F.data == "start_train_cb")
async def start_train_from_dict(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    cards = get_all_cards(user_id, shuffle=True)
    USER_TRAIN_STATE[user_id] = {
        'queue': cards,
        'index': 0,
        'session_count': 0
    }
    await send_train_card(callback, user_id, edit=True)

@router.callback_query(F.data == "export_excel")
async def export_excel_callback(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    cards = get_all_cards(user_id)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Словарь"
    ws.append(["№", "Фраза на английском", "Перевод на русский", "Статус", "Категория"])
    
    for idx, c in enumerate(cards, 1):
        st = "Знаю" if c.get('status') == 'known' else "Учу"
        ws.append([idx, c['phrase_en'], c['phrase_ru'], st, c.get('category', 'Общее')])
        
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    
    file = BufferedInputFile(bio.read(), filename="My_Vocabulary.xlsx")
    await callback.message.answer_document(file, caption="📄 Твой актуальный словарь в Excel со статусами!")
    await callback.answer()

@router.callback_query(F.data.startswith("tts_"))
async def tts_callback(callback: types.CallbackQuery):
    card_id = int(callback.data.split("_")[1])
    cards = get_all_cards(callback.from_user.id)
    card = next((c for c in cards if c['id'] == card_id), None)
    
    if not card or not card['phrase_en']:
        await callback.answer("Карточка не найдена", show_alert=True)
        return
        
    try:
        tts = gTTS(text=card['phrase_en'], lang='en', slow=False)
        bio = io.BytesIO()
        tts.write_to_fp(bio)
        bio.seek(0)
        audio_file = BufferedInputFile(bio.read(), filename=f"pronunciation_{card_id}.mp3")
        await callback.message.answer_voice(audio_file, caption=f"🗣️ Произношение: **{card['phrase_en']}**", parse_mode="Markdown")
        await callback.answer()
    except Exception as e:
        await callback.answer(f"Ошибка озвучки: {e}", show_alert=True)

# Document handler (Excel upload)
@router.message(F.document)
async def handle_document_upload(message: types.Message, bot: Bot):
    doc = message.document
    if not (doc.file_name.endswith('.xlsx') or doc.file_name.endswith('.xls')):
        await message.answer("Пожалуйста, отправьте Excel файл (`.xlsx`).")
        return
        
    file_info = await bot.get_file(doc.file_id)
    file_bytes = await bot.download_file(file_info.file_path)
    
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes.read()), data_only=True)
    ws = wb.active
    
    added_count = 0
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row_idx == 1:
            vals_str = " ".join([str(v or '') for v in row]).lower()
            if 'предложение' in vals_str or 'phrase' in vals_str or 'english' in vals_str:
                continue
                
        non_empty = [str(v).strip() for v in row if v is not None and str(v).strip()]
        if len(non_empty) >= 2:
            if non_empty[0].isdigit() and len(non_empty) >= 3:
                en = non_empty[1]
                ru = non_empty[2]
            else:
                en = non_empty[0]
                ru = non_empty[1]
                
            cat = categorize(en, ru)
            if add_card(message.from_user.id, en, ru, cat):
                added_count += 1
                
    kb = get_update_keyboard(message.from_user.id, is_single=False, added_count=added_count)
    await message.answer(
        f"🎉 **Отлично!** Успешно импортировано `{added_count}` карточек из Excel файла `{doc.file_name}`.\n\n"
        f"Они уже доступны в твоём словаре, группах и интерактивном тренажёре со статусом «Учу»!",
        reply_markup=kb,
        parse_mode="Markdown"
    )

SYSTEM_MENU_BUTTONS = {
    "🃏 Тренировка карточек", "📚 Мой словарь", "🗣️ Разговорный тренажёр", 
    "💼 Собеседование", "📊 Прогресс", "⚙️ Настройки", "📱 Открыть Mini App", "📱 Mini App"
}

@router.message(F.text == "/train")
async def cmd_train(message: types.Message):
    await start_flashcard_session(message)

# Text & Forward Handler
@router.message(F.text & ~F.text.startswith("/"))
async def handle_text_or_forward(message: types.Message):
    user_id = message.from_user.id
    text = message.text.strip()
    
    # Ignore system buttons
    if text in SYSTEM_MENU_BUTTONS:
        return
        
    # Check if user is in speaking or interview mode
    try:
        from handlers.speaking import USER_SPEAKING_MODE
        if USER_SPEAKING_MODE.get(user_id, False):
            return
    except Exception:
        pass
        
    try:
        from handlers.interview import USER_INTERVIEW_STATE
        if user_id in USER_INTERVIEW_STATE:
            return
    except Exception:
        pass

    u = get_or_create_user(user_id, message.from_user.username or "", message.from_user.full_name or "")
    t_lang = u.get('target_lang', 'en') if u else 'en'
    n_lang = u.get('native_lang', 'ru') if u else 'ru'
    
    t_meta = SUPPORTED_LANGS.get(t_lang, SUPPORTED_LANGS['en'])
    n_meta = SUPPORTED_LANGS.get(n_lang, SUPPORTED_LANGS['ru'])
    
    lines = [l.strip() for l in text.split('\n') if l.strip()]
    if len(lines) > 1:
        added = 0
        for l in lines:
            en, ru = parse_line_to_pair(l)
            if not en:
                continue
            if not ru:
                try:
                    tr = smart_translate(en, source_field="auto", native_lang=n_lang, target_lang=t_lang)
                    en = tr.get('text_target') or en
                    ru = tr.get('text_native') or ru
                    cat = tr.get('category') or categorize(en, ru)
                except Exception:
                    cat = categorize(en, ru)
            else:
                cat = categorize(en, ru)
            if add_card(user_id, en, ru, cat, target_lang=t_lang, native_lang=n_lang):
                added += 1
        kb = get_update_keyboard(user_id, is_single=False, added_count=added)
        await message.answer(
            f"✅ **Успешно добавлено `{added}` новых фраз в словарь и тренажёр со статусом «Учу»!**\n\n"
            f"Языковая пара: {n_meta['flag']} {n_meta['short']} ➔ {t_meta['flag']} {t_meta['short']}\n\n"
            f"Нажмите кнопку ниже, чтобы открыть обновлённый тренажёр или начать тренировку в чате:",
            reply_markup=kb,
            parse_mode="Markdown"
        )
        return
        
    en, ru = parse_line_to_pair(text)
    if not en:
        return
        
    if not ru:
        try:
            tr = smart_translate(en, source_field="auto", native_lang=n_lang, target_lang=t_lang)
            en = tr.get('text_target') or en
            ru = tr.get('text_native') or ru
            cat = tr.get('category') or categorize(en, ru)
        except Exception:
            cat = categorize(en, ru)
    else:
        cat = categorize(en, ru)
        
    card_id = add_card(user_id, en, ru, cat, target_lang=t_lang, native_lang=n_lang)
    
    kb = get_update_keyboard(user_id, card_id=card_id, is_single=True, added_count=1)
    
    reply_text = (
        f"🃏 **Карточка добавлена в словарь и тренажёр!**\n\n"
        f"{t_meta['flag']} **{en}**\n"
        f"{n_meta['flag']} {ru if ru else '*(без перевода)*'}\n"
        f"🏷 *Категория: {cat}* | 📌 *Статус: ⏳ Учу*\n\n"
        f"👇 *Используйте кнопки ниже для тренировки и обновления:*"
    )
    
    await message.answer(reply_text, reply_markup=kb, parse_mode="Markdown")

@router.callback_query(F.data.startswith("del_"))
async def delete_card_callback(callback: types.CallbackQuery):
    card_id = int(callback.data.split("_")[1])
    if delete_card(card_id, callback.from_user.id):
        await callback.message.edit_text("🗑️ Карточка удалена из словаря.")
        await callback.answer("Удалено")
    else:
        await callback.answer("Ошибка при удалении", show_alert=True)
